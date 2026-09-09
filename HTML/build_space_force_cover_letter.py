"""Generate a tailored cover letter for USAJOBS announcement 883267300."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

OUTPUT_DIR = Path(__file__).parent / "public"
DOCX_PATH = OUTPUT_DIR / "Jonathan_Witcoski_Cover_Letter_USSF_883267300.docx"
PDF_PATH = OUTPUT_DIR / "Jonathan_Witcoski_Cover_Letter_USSF_883267300.pdf"


def set_run_font(run, size=11, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_paragraph(doc, text="", bold=False, align=None, space_after=8):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    paragraph.paragraph_format.space_after = Pt(space_after)
    if text:
        run = paragraph.add_run(text)
        set_run_font(run, bold=bold)
    return paragraph


def build_document() -> Path:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    add_paragraph(doc, "Jonathan Witcoski", bold=True, space_after=2)
    add_paragraph(doc, "jwitcoski@gmail.com | (570) 582-3933", space_after=2)
    add_paragraph(doc, "Washington, DC metropolitan area (Reston, Virginia)", space_after=12)

    add_paragraph(doc, "September 9, 2026", space_after=12)
    add_paragraph(doc, "Hiring Manager", space_after=2)
    add_paragraph(doc, "United States Space Force", space_after=2)
    add_paragraph(doc, "Space Systems Command, Systems Delta 81", space_after=2)
    add_paragraph(doc, "Pentagon, Arlington, VA", space_after=12)

    subject = add_paragraph(doc, space_after=12)
    subject_run = subject.add_run(
        "Re: Chief Engineer, SYD 81 — Announcement SF-FY26-6S-NH-0801-4-SPLASH "
        "(Control Number 883267300)"
    )
    set_run_font(subject_run, bold=True)

    add_paragraph(doc, "Dear Hiring Manager:", space_after=8)

    paragraphs = [
        (
            "I am applying for the Chief Engineer, SYD 81 position with the United States Space Force. "
            "I bring more than 15 years of experience delivering enterprise geospatial applications, "
            "cloud-hosted data solutions, spatial databases, and automated software workflows for federal "
            "agencies and regulated organizations. I am drawn to this opportunity to help establish the "
            "systems engineering, integration, and technical governance needed to modernize secure, "
            "interoperable HR capabilities for the Space Force."
        ),
        (
            "In my current role as a Geospatial Engineer at INCATech, I design and maintain enterprise "
            "geodatabases and geospatial workflows supporting Intelligence Community and U.S. Postal "
            "Inspection Service law-enforcement missions. I automate data validation, spatial analysis, "
            "and extract-transform-load workflows using Python and SQL across Oracle, SQL Server, and "
            "PostgreSQL. This work requires managing system dependencies, translating mission needs into "
            "technical solutions, and delivering reliable capabilities within security-conscious federal "
            "environments."
        ),
        (
            "My broader experience aligns with the position's focus on cloud data, software delivery, and "
            "operational sustainment. At DRT Strategies, I built Python and R workflows producing PMTiles "
            "and GeoParquet datasets for public health dashboards and web maps, while developing SQL Server "
            "and PostgreSQL/PostGIS databases and recurring ETL pipelines. At the U.S. Census Bureau, I "
            "developed Python and ArcGIS API tools for large national datasets. As a GIS Data Engineer and "
            "Scrum Master at Saicon, I led Agile delivery of a cloud-hosted ArcGIS Enterprise migration and "
            "coordinated spatial data quality and milestone decisions across stakeholders."
        ),
        (
            "I am comfortable connecting architecture, data, interfaces, security considerations, testing, "
            "and user outcomes across multidisciplinary teams. I communicate technical tradeoffs clearly, "
            "work effectively with distributed stakeholders, and bring a practical, delivery-focused approach "
            "to improving system reliability and maintainability. My AWS Certified Cloud Practitioner "
            "credential and hands-on experience with AWS, Azure, Docker, Jenkins, GitLab, and CI/CD further "
            "support my ability to contribute to enterprise cloud modernization."
        ),
        (
            "Thank you for considering my application. I would welcome the opportunity to discuss how my "
            "federal mission experience, systems thinking, and record of delivering integrated data and "
            "software capabilities could support Systems Delta 81 and the United States Space Force."
        ),
    ]

    for text in paragraphs:
        add_paragraph(doc, text, space_after=8)

    add_paragraph(doc, "Sincerely,", space_after=18)
    add_paragraph(doc, "Jonathan Witcoski", bold=True, space_after=2)
    add_paragraph(doc, "witcoskitech.com", space_after=2)
    add_paragraph(doc, "Vector Scope AI LLC | vectorscopeai.com", space_after=0)

    doc.save(DOCX_PATH)
    return DOCX_PATH


def export_pdf(docx_path: Path, pdf_path: Path) -> int:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    document = None
    try:
        document = word.Documents.Open(str(docx_path.resolve()))
        pages = document.ComputeStatistics(2)
        document.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        return pages
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()


def main():
    docx_path = build_document()
    print(f"Saved Word: {docx_path}")

    try:
        pages = export_pdf(docx_path, PDF_PATH)
        print(f"Saved PDF: {PDF_PATH} ({pages} page(s))")
    except Exception as exc:
        print(f"PDF export skipped: {exc}")
        print("Open the Word file and save as PDF before submitting.")


if __name__ == "__main__":
    main()