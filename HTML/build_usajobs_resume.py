"""Generate a USAJOBS-compliant 2-page federal resume (Word + PDF)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

OUTPUT_DIR = Path(__file__).parent / "public"
DOCX_PATH = OUTPUT_DIR / "Jonathan_Witcoski_Resume_USAJOBS_2026.docx"
PDF_PATH = OUTPUT_DIR / "Jonathan_Witcoski_Resume_USAJOBS_2026.pdf"


def set_run_font(run, size=10, bold=False, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic


def add_section_header(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "666666")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_body(doc, text, bold=False, italic=False, space_after=2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run_font(run, bold=bold, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    run = p.add_run(text)
    set_run_font(run)
    return p


def add_job_block(doc, title, meta_lines, bullets):
    add_body(doc, title, bold=True, space_after=1)
    for line in meta_lines:
        add_body(doc, line, space_after=1)
    for bullet in bullets:
        add_bullet(doc, bullet)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    name.paragraph_format.space_after = Pt(2)
    run = name.add_run("JONATHAN WITCOSKI")
    set_run_font(run, size=14, bold=True)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact.paragraph_format.space_after = Pt(2)
    run = contact.add_run(
        "Email: jwitcoski@gmail.com | Phone: (570) 582-3933\n"
        "Washington, DC metropolitan area (previously Reston, Virginia) | Citizenship: United States"
    )
    set_run_font(run, size=10)

    add_section_header(doc, "Summary of Qualifications")
    add_body(
        doc,
        "Geographic Information Systems (GIS) developer with more than 15 years of experience designing, "
        "building, and supporting enterprise geospatial applications for federal agencies and regulated industries. "
        "Skilled in Esri ArcGIS Enterprise, ArcGIS Online, ArcGIS Pro, Python, R, JavaScript, Structured Query "
        "Language (SQL), SQL Server, PostgreSQL/PostGIS, REST application programming interfaces (APIs), and "
        "automated extract-transform-load (ETL) workflows. Experience includes public health surveillance mapping, "
        "utility GIS migration, census geography production, and emergency management support for the Department "
        "of Homeland Security (DHS) and the Federal Emergency Management Agency (FEMA).",
        space_after=4,
    )

    add_section_header(doc, "Work Experience")

    add_job_block(
        doc,
        "Geographer",
        [
            "Employer: DRT Strategies (Centers for Disease Control and Prevention) | Location: Remote",
            "Dates: 03/2022 – Present | Hours per week: 40 | Supervisor: Available upon request | May contact: Yes",
        ],
        [
            "Develop web mapping applications and dashboards using ArcGIS Enterprise, JavaScript, React, and Power BI for Centers for Disease Control and Prevention (CDC) public health surveillance programs.",
            "Use R and Python to build PMTiles and GeoParquet datasets that feed dashboards and web maps.",
            "Build SQL Server and PostgreSQL/PostGIS databases and Python/SQL ETL pipelines for recurring geospatial updates.",
        ],
    )

    add_job_block(
        doc,
        "Founder",
        [
            "Employer: Vector Scope AI LLC | Location: Remote",
            "Dates: 10/2025 – Present | Hours per week: 15 | Supervisor: Self | May contact: Yes",
        ],
        [
            "Build ArcGIS-connected web tools and Python pipelines for map editing, validation, and data extraction.",
            "Deliver Global Ski Atlas (3,000+ ski areas) as a public web mapping application with automated ETL.",
        ],
    )

    add_job_block(
        doc,
        "GIS Data Engineer and Scrum Master",
        [
            "Employer: Saicon (National Grid) | Location: Remote",
            "Dates: 05/2021 – 03/2022 | Hours per week: 40 | Supervisor: Available upon request | May contact: Yes",
        ],
        [
            "Migrated large utility GIS datasets to cloud-hosted ArcGIS Enterprise using Python and SQL.",
            "Led Agile delivery of migration milestones and spatial data quality checks.",
        ],
    )

    add_job_block(
        doc,
        "Geographer",
        [
            "Employer: U.S. Department of Commerce, U.S. Census Bureau | Location: Suitland, Maryland",
            "Dates: 11/2016 – 03/2021 | Hours per week: 40 | Series/Grade: Available upon request",
            "Supervisor: Available upon request | May contact: Yes",
        ],
        [
            "Developed Python tools and ArcGIS API workflows for census geography production.",
            "Automated spatial processing workflows across large national datasets.",
        ],
    )

    add_job_block(
        doc,
        "GIS Systems Administrator and Software Engineer",
        [
            "Employer: C2 Solutions Group Inc. | Location: Reston, Virginia",
            "Dates: 05/2014 – 11/2016 | Hours per week: 40 | Supervisor: Available upon request | May contact: Yes",
        ],
        [
            "Administered ArcGIS Server and Portal deployments backed by SQL Server for federal clients.",
            "Built secure web mapping applications used by hundreds of concurrent users.",
        ],
    )

    add_job_block(
        doc,
        "Geospatial Analyst",
        [
            "Employer: Booz Allen Hamilton (DHS/FEMA) | Location: Philadelphia, Pennsylvania and Arlington, Virginia",
            "Dates: 10/2009 – 05/2014 | Hours per week: 40 | Supervisor: Available upon request | May contact: Yes",
        ],
        [
            "Built maps, spatial analysis workflows, and web GIS tools for FEMA and DHS disaster response operations.",
            "Supported situational awareness and decision-making during major disaster events.",
        ],
    )

    add_job_block(
        doc,
        "Earlier GIS Roles (2007 – 2009)",
        [
            "Employer: Tyco Telecommunications; WDG Location Consulting; Philmont Scout Ranch",
            "Dates: 05/2007 – 05/2009 | Hours per week: 40 (full-time and seasonal roles)",
            "Supervisor: Available upon request | May contact: Yes",
        ],
        [
            "GIS solutions development, location consulting analysis, and field GIS support.",
        ],
    )

    add_section_header(doc, "Education")
    add_body(doc, "Master of Science, Geography | University of Tennessee, Knoxville | 05/2007", space_after=1)
    add_body(
        doc,
        "Bachelor of Arts, Geography and Anthropology; Minor in Geographic Information Systems | Pennsylvania State University | 05/2005",
        space_after=4,
    )

    add_section_header(doc, "Certifications")
    add_body(doc, "AWS Certified Cloud Practitioner | Amazon Web Services | 03/2021", space_after=4)

    add_section_header(doc, "Technical Skills")
    add_body(
        doc,
        "ArcGIS Enterprise, ArcGIS Online, ArcGIS Pro, ArcGIS API for JavaScript, Spatial Analyst, Python, R, "
        "JavaScript, SQL, SQL Server, PostgreSQL/PostGIS, GeoParquet, PMTiles, REST APIs, web services, AWS, Azure, "
        "Docker, Jenkins, GitLab, continuous integration/continuous delivery (CI/CD), React, Power BI, Agile/Scrum",
        space_after=0,
    )

    doc.save(DOCX_PATH)
    return DOCX_PATH


def export_pdf(docx_path: Path, pdf_path: Path) -> int:
    import win32com.client

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    try:
        document = word.Documents.Open(str(docx_path.resolve()))
        pages = document.ComputeStatistics(2)
        document.SaveAs(str(pdf_path.resolve()), FileFormat=17)
        document.Close(False)
        return pages
    finally:
        word.Quit()


def main():
    docx_path = build_document()
    print(f"Saved Word: {docx_path}")

    try:
        pages = export_pdf(docx_path, PDF_PATH)
        print(f"Saved PDF: {PDF_PATH} ({pages} page(s))")
        if pages > 2:
            print("WARNING: Resume exceeds 2 pages. Trim content before uploading to USAJOBS.")
    except Exception as exc:
        print(f"PDF export skipped: {exc}")
        print("Open the Word file and save as PDF before uploading to USAJOBS.")


if __name__ == "__main__":
    main()
