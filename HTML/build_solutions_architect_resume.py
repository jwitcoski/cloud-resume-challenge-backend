"""Generate Solutions Architect–oriented resume Word document."""
import re
import shutil
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

INPUT = Path(__file__).parent / "public" / "Jonathan_Witcoski_Resume_Solutions_Architect_2026.md"
OUTPUT = Path(__file__).parent / "public" / "Jonathan_Witcoski_Resume_Solutions_Architect_2026.docx"
DOWNLOADS = Path.home() / "Downloads" / "Jonathan_Witcoski_Resume_Solutions_Architect_2026.docx"


def set_run_font(run, size=11, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold


def parse_bold_segments(paragraph, text, size=11):
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    parts = re.split(r"(\*\*[^*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        i += 1
        if not line:
            continue
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[2:])
            set_run_font(run, size=16, bold=True)
            continue
        if line.startswith("## "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].upper())
            set_run_font(run, size=11, bold=True)
            continue
        if line.startswith("### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[4:])
            set_run_font(run, size=11, bold=True)
            continue
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(2)
            parse_bold_segments(p, line[2:])
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        parse_bold_segments(p, line)

    doc.save(docx_path)


if __name__ == "__main__":
    convert(INPUT, OUTPUT)
    shutil.copy2(OUTPUT, DOWNLOADS)
    print(f"Saved: {OUTPUT}")
    print(f"Copied: {DOWNLOADS}")
